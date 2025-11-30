from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from datetime import datetime

# File path
xlsx_path = r'C:\Users\mwaelterman\OneDrive - WRB\Mark\Data\Excel\Personal\2023-2024 FastCAT Results.xlsx'

print(f"Loading workbook: {xlsx_path}")
print("This may take a while for large files...")

# Load workbook preserving formulas (data_only=False is default)
wb = load_workbook(xlsx_path)
ws = wb.active

total_rows = ws.max_row
total_cols = ws.max_column
print(f"Loaded. Total rows: {total_rows:,}, columns: {total_cols}")

# Create new workbook and copy only rows to KEEP (where column D year != 2025)
print("Filtering rows (copying rows to keep)...")
new_wb = Workbook()
new_ws = new_wb.active
new_ws.title = ws.title

# Copy column dimensions
for col in range(1, total_cols + 1):
    col_letter = get_column_letter(col)
    if ws.column_dimensions[col_letter].width:
        new_ws.column_dimensions[col_letter].width = ws.column_dimensions[col_letter].width

rows_kept = 0
rows_deleted = 0

for row_num in range(1, total_rows + 1):
    # Always keep header row
    if row_num == 1:
        for col in range(1, total_cols + 1):
            src_cell = ws.cell(row=row_num, column=col)
            dst_cell = new_ws.cell(row=rows_kept + 1, column=col)
            dst_cell.value = src_cell.value
            if src_cell.has_style:
                dst_cell.font = src_cell.font.copy()
                dst_cell.fill = src_cell.fill.copy()
                dst_cell.border = src_cell.border.copy()
                dst_cell.alignment = src_cell.alignment.copy()
                dst_cell.number_format = src_cell.number_format
        rows_kept += 1
        continue
    
    cell_d = ws.cell(row=row_num, column=4)
    cell_value = cell_d.value
    
    # Check if this row should be deleted (year = 2025)
    should_delete = False
    if isinstance(cell_value, datetime):
        if cell_value.year == 2025:
            should_delete = True
    elif isinstance(cell_value, (int, float)):
        if int(cell_value) == 2025:
            should_delete = True
    
    if should_delete:
        rows_deleted += 1
    else:
        # Copy row to new workbook
        rows_kept += 1
        for col in range(1, total_cols + 1):
            src_cell = ws.cell(row=row_num, column=col)
            dst_cell = new_ws.cell(row=rows_kept, column=col)
            dst_cell.value = src_cell.value
            if src_cell.has_style:
                dst_cell.font = src_cell.font.copy()
                dst_cell.fill = src_cell.fill.copy()
                dst_cell.border = src_cell.border.copy()
                dst_cell.alignment = src_cell.alignment.copy()
                dst_cell.number_format = src_cell.number_format
    
    if row_num % 50000 == 0:
        print(f"Processed {row_num:,} of {total_rows:,} rows... (kept: {rows_kept:,}, deleted: {rows_deleted:,})")

print(f"Deleted {rows_deleted:,} rows, kept {rows_kept:,} rows")

print("Saving...")
wb.close()
new_wb.save(xlsx_path)
new_wb.close()

print('Filtering complete!')
